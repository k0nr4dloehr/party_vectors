"""Preview-render a PPTX slide to PNG for visual verification.

Reads the real .pptx (shapes, positions, fills, text) and redraws with
matplotlib so we can eyeball the layout without LibreOffice.  Not for
production; a developer aid only.
"""

from __future__ import annotations

import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrow, FancyBboxPatch, Rectangle
from pptx import Presentation
from pptx.util import Emu

SRC = Path(sys.argv[1] if len(sys.argv) > 1 else "figures/paper/methodology_pipeline.pptx")
PAGE_W_IN, PAGE_H_IN = 10.0, 7.5


def _in(emu) -> float:
    return Emu(emu).inches


def main() -> None:
    prs = Presentation(SRC)
    for si, slide in enumerate(prs.slides):
        fig = plt.figure(figsize=(PAGE_W_IN, PAGE_H_IN), dpi=150)
        ax = fig.add_axes([0, 0, 1, 1])
        ax.set_xlim(0, PAGE_W_IN)
        ax.set_ylim(0, PAGE_H_IN)
        ax.invert_yaxis()
        ax.axis("off")
        for sh in slide.shapes:
            x = _in(sh.left)
            y = _in(sh.top)
            w = _in(sh.width)
            h = _in(sh.height)
            st = sh.shape_type
            if str(st) == "AUTO_SHAPE (1)":
                fill = None
                try:
                    fill = sh.fill.fore_color.rgb
                except Exception:
                    pass
                lc = None
                try:
                    lc = sh.line.color.rgb
                except Exception:
                    pass
                fc = "#" + str(fill) if fill else "none"
                ec = "#" + str(lc) if lc else "0.4"
                box = FancyBboxPatch(
                    (x, y),
                    w,
                    h,
                    boxstyle="round,pad=0.02,rounding_size=0.03",
                    fc=fc,
                    ec=ec,
                    lw=1.1,
                )
                ax.add_patch(box)
                # center text
                try:
                    tx = sh.text_frame.text
                except Exception:
                    tx = ""
                if tx.strip():
                    ax.text(
                        x + w / 2,
                        y + h / 2,
                        tx,
                        ha="center",
                        va="center",
                        fontsize=6,
                        wrap=True,
                    )
            elif str(st) == "LINE (9)":
                lc = "0.3"
                try:
                    lc = "#" + str(sh.line.color.rgb)
                except Exception:
                    pass
                ax.add_patch(
                    FancyArrow(
                        x,
                        y,
                        x + w,
                        y + h,
                        width=0.012,
                        head_width=0.05,
                        head_length=0.09,
                        fc=lc,
                        ec=lc,
                        length_includes_head=True,
                    )
                )
            else:  # TEXT_BOX
                tx = sh.text_frame.text
                if not tx.strip():
                    continue
                ax.text(
                    x + w / 2,
                    y + h / 2,
                    tx,
                    ha="center",
                    va="center",
                    fontsize=6,
                )
        out = SRC.with_name(SRC.stem + f"_preview{si}.png")
        fig.savefig(out)
        plt.close(fig)
        print("wrote", out)


if __name__ == "__main__":
    main()